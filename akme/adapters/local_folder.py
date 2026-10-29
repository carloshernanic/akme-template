from __future__ import annotations

import csv
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

from .base import Adapter, AdapterError, require

TABLES = {".csv", ".xlsx"}
DOCS = {".md", ".txt", ".pdf", ".docx", ".pptx"}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
OFFICE_MISSING = "[{} ignorado: instale o extra 'office' (uv sync --extra office) para extrair texto]"


class LocalFolderAdapter(Adapter):
    name = "local_folder"

    def __init__(self, path: str | list[str] = "data", encoding: str = "utf-8", delimiter: str = ",",
                 recursive: bool = True) -> None:
        self.roots = [Path(p) for p in ([path] if isinstance(path, str) else path)]
        missing = [str(r) for r in self.roots if not r.is_dir()]
        if missing:
            raise AdapterError(f"pasta nao encontrada: {', '.join(missing)}")
        self.encoding = encoding
        self.delimiter = delimiter
        self.recursive = recursive
        self._tables: dict[str, tuple[Path, Path, str | None]] | None = None
        self._docs: dict[str, Path] | None = None

    def _walk(self, suffixes: set[str]) -> list[tuple[Path, Path]]:
        found = []
        for root in self.roots:
            for file in sorted(root.rglob("*") if self.recursive else root.iterdir()):
                rel = file.relative_to(root)
                if (file.is_file() and file.suffix.lower() in suffixes and not file.name.startswith("~$")
                        and not any(part.startswith(".") for part in rel.parts)):
                    found.append((root, file))
        return found

    def _index(self) -> dict[str, tuple[Path, Path, str | None]]:
        if self._tables is None:
            index: dict[str, tuple[Path, Path, str | None]] = {}
            for root, file in self._walk(TABLES):
                parts = [*file.relative_to(root).parent.parts, file.stem]
                sheets = _sheet_names(file) if file.suffix.lower() == ".xlsx" else [None]
                for sheet in sheets:
                    name_parts = parts + ([sheet] if len(sheets) > 1 else [])
                    _register(index, "__".join(_slug(p) for p in name_parts), root, (root, file, sheet))
            self._tables = index
        return self._tables

    def _doc_index(self) -> dict[str, Path]:
        if self._docs is None:
            index: dict[str, Path] = {}
            for root, file in self._walk(DOCS):
                _register(index, file.relative_to(root).as_posix(), root, file, sep="/")
            self._docs = index
        return self._docs

    def _source(self, source: str) -> tuple[Path, Path, str | None]:
        index = self._index()
        if source not in index:
            raise AdapterError(f"fonte nao encontrada: {source}. Fontes: {', '.join(sorted(index)) or 'nenhuma'}")
        return index[source]

    def list_sources(self) -> list[str]:
        return sorted(self._index())

    def source_location(self, source: str) -> dict[str, Any]:
        root, file, sheet = self._source(source)
        return {"file": file, "path": file.relative_to(root).as_posix(), "sheet": sheet}

    def _table(self, source: str, n: int) -> tuple[list[str], list[list[Any]]]:
        _, file, sheet = self._source(source)
        if file.suffix.lower() == ".xlsx":
            header, rows = _xlsx_table(file, sheet, n)
        else:
            with file.open(encoding=self.encoding, newline="") as fh:
                reader = csv.reader(fh, delimiter=self.delimiter)
                header = next(reader, [])
                rows = []
                for row in reader:
                    if len(rows) >= n:
                        break
                    rows.append([v if v != "" else None for v in row])
        columns = _header(header)
        return columns, [(r + [None] * len(columns))[: len(columns)] for r in rows]

    def read_schema(self, source: str) -> dict[str, Any]:
        columns, _ = self._table(source, 0)
        return {"name": source, "columns": [{"name": c, "type": "string"} for c in columns]}

    def read_sample(self, source: str, n: int = 100) -> list[dict[str, Any]]:
        columns, rows = self._table(source, n)
        return [dict(zip(columns, r)) for r in rows]

    def _doc(self, name: str, file: Path) -> dict[str, Any]:
        pages = _doc_pages(file, self.encoding)
        text = pages[0]["text"] if len(pages) == 1 else "\n\n".join(f"[{p['label']}]\n{p['text']}" for p in pages)
        return {"name": name, "path": file.as_posix(), "text": text, "pages": pages}

    def read_docs(self) -> list[dict[str, Any]]:
        return [self._doc(name, file) for name, file in self._doc_index().items()]

    def read_doc(self, name: str) -> dict[str, Any] | None:
        file = self._doc_index().get(name)
        return self._doc(name, file) if file else None


def _register(index: dict, name: str, root: Path, value: Any, sep: str = "__") -> None:
    if name in index:
        name = f"{_slug(root.name) if sep == '__' else root.name}{sep}{name}"
    if name in index:
        raise AdapterError(f"dois arquivos geram o mesmo nome '{name}'; renomeie um deles")
    index[name] = value


def _slug(text: str) -> str:
    ascii_text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    return re.sub(r"[^0-9a-zA-Z]+", "_", ascii_text).strip("_").lower() or "x"


def _header(values: list[Any]) -> list[str]:
    names: list[str] = []
    for i, v in enumerate(values, 1):
        base = str(v).strip() if v not in (None, "") else f"col_{i}"
        name, k = base, 2
        while name in names:
            name, k = f"{base}_{k}", k + 1
        names.append(name)
    return names


def _cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip() or None
    return value


def _sheet_names(file: Path) -> list[str]:
    require("openpyxl", "office")
    import openpyxl
    wb = openpyxl.load_workbook(file, read_only=True)
    try:
        return [ws.title for ws in wb.worksheets]
    finally:
        wb.close()


def _xlsx_table(file: Path, sheet: str | None, n: int) -> tuple[list[Any], list[list[Any]]]:
    require("openpyxl", "office")
    import openpyxl
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    try:
        header: list[Any] | None = None
        rows: list[list[Any]] = []
        for raw in wb[sheet].iter_rows(values_only=True):
            values = [_cell(v) for v in raw]
            if all(v is None for v in values):
                continue
            if header is None:
                header = values
                while header and header[-1] is None:
                    header.pop()
                continue
            if len(rows) >= n:
                break
            rows.append(values)
        return header or [], rows
    finally:
        wb.close()


def _row_text(cells: list[str]) -> str:
    out: list[str] = []
    for c in (c.strip() for c in cells):
        if not out or c != out[-1]:
            out.append(c)
    return " | ".join(out)


def _doc_pages(file: Path, encoding: str) -> list[dict[str, str]]:
    suffix = file.suffix.lower()
    try:
        if suffix == ".pdf":
            texts = _pdf_pages(file)
        elif suffix == ".docx":
            texts = _docx_pages(file)
        elif suffix == ".pptx":
            return [{"label": f"slide {i}", "text": t} for i, t in enumerate(_pptx_pages(file), 1)]
        else:
            texts = [file.read_text(encoding=encoding)]
    except ImportError:
        texts = [OFFICE_MISSING.format(suffix[1:]) if suffix != ".pdf"
                 else "[pdf ignorado: instale o extra 'pdf' para extrair texto]"]
    except Exception as exc:
        texts = [f"[erro ao ler {file.name}: {exc}]"]
    return [{"label": f"pagina {i}", "text": t} for i, t in enumerate(texts or [""], 1)]


def _pdf_pages(file: Path) -> list[str]:
    import pypdf
    return [page.extract_text() or "" for page in pypdf.PdfReader(str(file)).pages]


def _docx_pages(file: Path) -> list[str]:
    from docx import Document
    from docx.table import Table
    doc = Document(str(file))
    pages: list[str] = []
    buf: list[str] = []

    def new_page() -> None:
        if any(x.strip() for x in buf):
            pages.append("\n".join(x for x in buf if x.strip()))
            buf.clear()

    for child in doc.element.body.iterchildren():
        if child.tag == W + "tbl":
            buf.extend(_row_text([c.text for c in row.cells]) for row in Table(child, doc).rows)
        elif child.tag == W + "p":
            line: list[str] = []
            for el in child.iter():
                if el.tag == W + "t":
                    line.append(el.text or "")
                elif el.tag == W + "tab":
                    line.append("\t")
                elif el.tag == W + "lastRenderedPageBreak" or (el.tag == W + "br" and el.get(W + "type") == "page"):
                    buf.append("".join(line))
                    line = []
                    new_page()
            buf.append("".join(line))
    new_page()
    return pages


def _pptx_pages(file: Path) -> list[str]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def shape_text(shape: Any) -> list[str]:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            return [t for s in shape.shapes for t in shape_text(s)]
        if getattr(shape, "has_table", False):
            return [_row_text([c.text for c in row.cells]) for row in shape.table.rows]
        if shape.has_text_frame:
            return [shape.text_frame.text]
        return []

    slides = []
    for slide in Presentation(str(file)).slides:
        lines = [t for shape in slide.shapes for t in shape_text(shape)]
        frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        if frame is not None and frame.text.strip():
            lines.append(f"Notas: {frame.text.strip()}")
        slides.append("\n".join(x for x in lines if x.strip()))
    return slides
